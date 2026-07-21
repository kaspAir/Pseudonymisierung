"""Erzeugt ein Word-Dokument aus einer Markdown-Datei dieses Repos.

Bewusst ein KONVERTER und keine zweite Fassung des Textes: eine von Hand
gepflegte Word-Datei liefe unweigerlich vom Markdown weg, und dann gaebe es
zwei Wahrheiten ueber dieselbe Schnittstelle.

Unterstuetzt genau die Auszeichnungen, die in den Dokumenten dieses Repos
vorkommen: Ueberschriften, Absaetze, Tabellen, Codebloecke, Zitatbloecke,
Aufzaehlungen, Trennlinien sowie inline `Code`, **fett** und *kursiv*.

    python scripts/md_nach_docx.py docs/API.md [ziel.docx]
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

CODE_SCHRIFT = "Consolas"
CODE_HINTERGRUND = "F2F2F2"
RAHMEN_GRAU = "BFBFBF"


# ------------------------------------------------------------- Hilfsmittel

def _schattierung(element, farbe):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), farbe)
    element.append(shd)


def _zellen_schattierung(zelle, farbe):
    _schattierung(zelle._tc.get_or_add_tcPr(), farbe)


def _absatz_schattierung(absatz, farbe):
    _schattierung(absatz._p.get_or_add_pPr(), farbe)


def _rahmen_unten(absatz, farbe=RAHMEN_GRAU, staerke=6):
    pPr = absatz._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    unten = OxmlElement("w:bottom")
    unten.set(qn("w:val"), "single")
    unten.set(qn("w:sz"), str(staerke))
    unten.set(qn("w:space"), "1")
    unten.set(qn("w:color"), farbe)
    pbdr.append(unten)
    pPr.append(pbdr)


_RE_INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def _schreibe_inline(absatz, text, fett=False, kursiv=False, unterstrichen=False):
    """Setzt `Code`, **fett**, *kursiv* und [Text](Ziel) in Laeufe um.

    Durchgaengig rekursiv, weil sich die Auszeichnungen im Text dieses Repos
    verschachteln: **fett mit `Code` darin**, [`API.md`](API.md). Ohne Rekursion
    blieben die inneren Zeichen sichtbar stehen.
    """
    for teil in _RE_INLINE.split(text):
        if not teil:
            continue
        if teil.startswith("`") and teil.endswith("`"):
            lauf = absatz.add_run(teil[1:-1])
            lauf.font.name = CODE_SCHRIFT
            lauf.font.size = Pt(9)
            lauf.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
            lauf.bold, lauf.italic, lauf.underline = fett, kursiv, unterstrichen
        elif teil.startswith("**") and teil.endswith("**"):
            _schreibe_inline(absatz, teil[2:-2], True, kursiv, unterstrichen)
        elif teil.startswith("*") and teil.endswith("*"):
            _schreibe_inline(absatz, teil[1:-1], fett, True, unterstrichen)
        elif teil.startswith("["):
            _schreibe_inline(absatz, teil[1:teil.index("](")], fett, kursiv, True)
        else:
            lauf = absatz.add_run(teil)
            lauf.bold, lauf.italic, lauf.underline = fett, kursiv, unterstrichen


# ------------------------------------------------------------------ Bloecke

def _nicht_trennen(zeile):
    """Verhindert, dass eine Tabellenzeile ueber den Seitenumbruch reisst."""
    trPr = zeile._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _kopfzeile_wiederholen(zeile):
    trPr = zeile._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def _tabelle(dok, zeilen):
    kopf = [z.strip() for z in zeilen[0].strip("|").split("|")]
    daten = []
    for zeile in zeilen[2:]:
        daten.append([z.strip() for z in zeile.strip("|").split("|")])

    # Markdown-Tabellen ohne Spaltentitel (»| | |«) erzeugen sonst eine leere,
    # grau hinterlegte Kopfzeile.
    hat_kopf = any(t for t in kopf)

    tabelle = dok.add_table(rows=1 if hat_kopf else 0, cols=len(kopf))
    tabelle.style = "Table Grid"
    tabelle.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabelle.autofit = True

    if hat_kopf:
        for i, titel in enumerate(kopf):
            zelle = tabelle.rows[0].cells[i]
            zelle.text = ""
            _schreibe_inline(zelle.paragraphs[0], titel)
            for lauf in zelle.paragraphs[0].runs:
                lauf.bold = True
            _zellen_schattierung(zelle, "E8E8E8")
        _nicht_trennen(tabelle.rows[0])
        _kopfzeile_wiederholen(tabelle.rows[0])

    for satz in daten:
        zeile = tabelle.add_row()
        _nicht_trennen(zeile)
        for i in range(len(kopf)):
            zelle = zeile.cells[i]
            zelle.text = ""
            wert = satz[i] if i < len(satz) else ""
            _schreibe_inline(zelle.paragraphs[0], wert.replace("\\|", "|"))
            for absatz in zelle.paragraphs:
                absatz.paragraph_format.space_after = Pt(2)
                for lauf in absatz.runs:
                    lauf.font.size = Pt(9)

    dok.add_paragraph().paragraph_format.space_after = Pt(6)


def _codeblock(dok, zeilen):
    for text in zeilen:
        absatz = dok.add_paragraph()
        absatz.paragraph_format.space_after = Pt(0)
        absatz.paragraph_format.space_before = Pt(0)
        absatz.paragraph_format.left_indent = Cm(0.4)
        _absatz_schattierung(absatz, CODE_HINTERGRUND)
        lauf = absatz.add_run(text if text else " ")
        lauf.font.name = CODE_SCHRIFT
        lauf.font.size = Pt(8.5)
    dok.add_paragraph().paragraph_format.space_after = Pt(6)


def wandle(quelle, ziel):
    with open(quelle, encoding="utf-8") as f:
        zeilen = f.read().splitlines()

    dok = Document()
    normal = dok.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for abschnitt in dok.sections:
        abschnitt.left_margin = Cm(2.2)
        abschnitt.right_margin = Cm(2.2)
        abschnitt.top_margin = Cm(2.0)
        abschnitt.bottom_margin = Cm(2.0)

    i = 0
    while i < len(zeilen):
        zeile = zeilen[i]

        # Codeblock
        if zeile.startswith("```"):
            j = i + 1
            block = []
            while j < len(zeilen) and not zeilen[j].startswith("```"):
                block.append(zeilen[j])
                j += 1
            _codeblock(dok, block)
            i = j + 1
            continue

        # Tabelle
        if zeile.startswith("|") and i + 1 < len(zeilen) and \
                re.match(r"^\|[\s:|-]+\|$", zeilen[i + 1].strip()):
            j = i
            block = []
            while j < len(zeilen) and zeilen[j].startswith("|"):
                block.append(zeilen[j])
                j += 1
            _tabelle(dok, block)
            i = j
            continue

        # Trennlinie
        if zeile.strip() == "---":
            absatz = dok.add_paragraph()
            _rahmen_unten(absatz)
            i += 1
            continue

        # Ueberschriften
        treffer = re.match(r"^(#{1,4})\s+(.*)$", zeile)
        if treffer:
            stufe = len(treffer.group(1))
            ueberschrift = dok.add_heading("", level=stufe)
            _schreibe_inline(ueberschrift, treffer.group(2))
            for lauf in ueberschrift.runs:
                lauf.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
            i += 1
            continue

        # Zitatblock
        if zeile.startswith("> "):
            j = i
            block = []
            while j < len(zeilen) and zeilen[j].startswith(">"):
                block.append(zeilen[j].lstrip(">").strip())
                j += 1
            absatz = dok.add_paragraph()
            absatz.paragraph_format.left_indent = Cm(0.6)
            _absatz_schattierung(absatz, "FFF6E0")
            _schreibe_inline(absatz, " ".join(t for t in block if t))
            i = j
            continue

        # Aufzaehlung - umgebrochene Folgezeilen gehoeren zum selben Punkt,
        # sonst reisst ein langer Punkt in einen eigenen Absatz ab.
        treffer = re.match(r"^\s*(?:[-*]|\d+\.)\s+(.*)$", zeile)
        if treffer:
            aufzaehlend = re.match(r"^\s*\d+\.", zeile) is not None
            block = [treffer.group(1)]
            j = i + 1
            while j < len(zeilen) and zeilen[j].strip() and \
                    not re.match(r"^(#|\||```|>|\s*[-*]\s|\s*\d+\.\s|---$)",
                                 zeilen[j]):
                block.append(zeilen[j].strip())
                j += 1
            absatz = dok.add_paragraph(
                style="List Number" if aufzaehlend else "List Bullet")
            absatz.paragraph_format.space_after = Pt(2)
            _schreibe_inline(absatz, " ".join(block))
            i = j
            continue

        # Leerzeile
        if not zeile.strip():
            i += 1
            continue

        # Fliesstext - Folgezeilen bis zur naechsten Leerzeile anhaengen
        block = [zeile]
        j = i + 1
        while j < len(zeilen) and zeilen[j].strip() and \
                not re.match(r"^(#|\||```|>|\s*[-*]\s|\s*\d+\.\s|---$)", zeilen[j]):
            block.append(zeilen[j])
            j += 1
        absatz = dok.add_paragraph()
        absatz.paragraph_format.space_after = Pt(6)
        _schreibe_inline(absatz, " ".join(t.strip() for t in block))
        i = j

    dok.save(ziel)
    return ziel


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    quelle = sys.argv[1]
    ziel = sys.argv[2] if len(sys.argv) > 2 else \
        os.path.splitext(quelle)[0] + ".docx"
    wandle(quelle, ziel)
    print("Geschrieben: %s" % ziel)
    return 0


if __name__ == "__main__":
    sys.exit(main())
