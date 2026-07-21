"""Erzeugt die Namenslexika aus den BfS-Dateien.

Quelle: Bundesamt fuer Statistik (BfS), Namen der staendigen Wohnbevoelkerung.
Die Dateien werden NICHT ins Repo uebernommen - nur die daraus abgeleiteten
Listen, damit die Herkunft nachvollziehbar bleibt.

Aufruf:
    python scripts/importiere_namen.py <Nachnamen.xlsx> <Vornamen.xlsx>
"""
import os
import sys

import openpyxl

ZIEL = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "lexikon")

# Ab welcher Zeile die Daten beginnen (Zeile 6 = Spaltenkoepfe).
ERSTE_DATENZEILE = 7


def _saubere(wert):
    if wert is None:
        return None
    s = str(wert).strip()
    return s or None


def lies_nachnamen(pfad):
    """Blatt 'CH': LASTNAME | RANK_CH | NUMBER_CH."""
    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    ws = wb["CH"]
    namen = {}
    for i, zeile in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < ERSTE_DATENZEILE:
            continue
        name = _saubere(zeile[0])
        if not name:
            continue
        try:
            anzahl = int(zeile[2])
        except (TypeError, ValueError):
            anzahl = 0
        namen[name] = max(namen.get(name, 0), anzahl)
    wb.close()
    return namen


def lies_vornamen(pfad, blatt=None):
    """Blatt '<Jahr>': Vorname | weiblich | maennlich ('*' = unterdrueckt)."""
    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    blatt = blatt or wb.sheetnames[0]
    ws = wb[blatt]
    namen = {}
    for i, zeile in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < ERSTE_DATENZEILE:
            continue
        name = _saubere(zeile[0])
        if not name:
            continue
        anzahl = 0
        for spalte in (1, 2):
            try:
                anzahl += int(zeile[spalte])
            except (TypeError, ValueError):
                pass          # '*' = aus Datenschutzgruenden unterdrueckt
        namen[name] = max(namen.get(name, 0), anzahl)
    wb.close()
    return namen, blatt


def schreibe(dateiname, namen, kopf):
    os.makedirs(ZIEL, exist_ok=True)
    pfad = os.path.join(ZIEL, dateiname)
    with open(pfad, "w", encoding="utf-8", newline="\n") as f:
        for zeile in kopf:
            f.write("# %s\n" % zeile)
        f.write("#\n")
        for name in sorted(namen):
            f.write("%s\n" % name)
    return pfad, len(namen)


# Ab wie vielen Traegern ein Name als alleinstehendes Signal taugt.
#
# Ohne verlaessliche Grossschreibung (Diktat) ist JEDES Wort ein Kandidat.
# Dann feuern Namen, die zwar in der Liste stehen, aber praktisch nie eine
# Person meinen: "Unser" hat 4 Traeger, "Server" 7 - gegenueber Buergi 2999
# und Vogt 4801. Die Schwelle trennt das, ohne etwas zu erfinden.
SCHWELLE_HAEUFIG = 20


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1

    nachnamen = lies_nachnamen(sys.argv[1])
    vornamen, blatt = lies_vornamen(sys.argv[2])

    pfad_n, anzahl_n = schreibe(
        "nachnamen.txt", nachnamen,
        ["Nachnamen der staendigen Wohnbevoelkerung der Schweiz",
         "Quelle: Bundesamt fuer Statistik (BfS), Blatt 'CH'",
         "Abgeleitet, nicht erfunden. %d Eintraege." % len(nachnamen)],
    )
    pfad_v, anzahl_v = schreibe(
        "vornamen.txt", vornamen,
        ["Vornamen der Bevoelkerung der Schweiz",
         "Quelle: Bundesamt fuer Statistik (BfS), Blatt '%s'" % blatt,
         "Abgeleitet, nicht erfunden. %d Eintraege." % len(vornamen)],
    )

    haeufig_n = {n: c for n, c in nachnamen.items() if c >= SCHWELLE_HAEUFIG}
    haeufig_v = {n: c for n, c in vornamen.items() if c >= SCHWELLE_HAEUFIG}
    pfad_hn, anzahl_hn = schreibe(
        "nachnamen_haeufig.txt", haeufig_n,
        ["Nachnamen mit mindestens %d Traegern" % SCHWELLE_HAEUFIG,
         "Quelle: Bundesamt fuer Statistik (BfS), Blatt 'CH'",
         "Nur fuer ALLEINSTEHENDE Treffer in Text ohne Grossschreibung.",
         "%d von %d Nachnamen." % (len(haeufig_n), len(nachnamen))],
    )
    pfad_hv, anzahl_hv = schreibe(
        "vornamen_haeufig.txt", haeufig_v,
        ["Vornamen mit mindestens %d Traegern" % SCHWELLE_HAEUFIG,
         "Quelle: Bundesamt fuer Statistik (BfS), Blatt '%s'" % blatt,
         "Nur fuer ALLEINSTEHENDE Treffer in Text ohne Grossschreibung.",
         "%d von %d Vornamen." % (len(haeufig_v), len(vornamen))],
    )

    print("Nachnamen: %6d -> %s" % (anzahl_n, pfad_n))
    print("Vornamen:  %6d -> %s" % (anzahl_v, pfad_v))
    print("davon haeufig (>= %d Traeger): %d Nachnamen, %d Vornamen"
          % (SCHWELLE_HAEUFIG, anzahl_hn, anzahl_hv))
    ueberschneidung = set(nachnamen) & set(vornamen)
    print("In beiden Listen: %d" % len(ueberschneidung))
    return 0


if __name__ == "__main__":
    sys.exit(main())
